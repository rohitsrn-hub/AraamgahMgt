"""
Excel export matching the guesthouse's manual "Summary of Room Occupancy"
paper ledger format, one row per room per booking for a given month.

Uses booking_financials() from report_calc.py — the same shared calculation
used by the room-allotment/guest-details/room-occupancy/monthly reports — so
this export can never drift from the app's other financial reports.

Columns Army No / Rank / Unit / Bill No are intentionally left blank: the
app does not store military rank, service number or unit (removed during
the sanitization pass as confidential, paper-only fields — see
generateOrgDataForm in pdfUtils.js) and "Bill No" is the guesthouse's
physical voucher-book number, which has no digital equivalent. Mobile
Number and Org/Non-Org are included instead so a row can still be matched
to a person by name without those fields — mobile number is captured on
every booking, unlike Aadhaar which has no capture point anywhere in the
app's UI today.
"""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.report_calc import booking_financials, build_room_maps, report_booking_query

# Column layout mirrors the manual sheet exactly, with two extra identification
# columns (Mobile No, Org/Non-Org) inserted after Name.
HEADERS = [
    "Ser\nNo", "Army No", "Rank", "Name", "Mobile No", "Org /\nNon-Org", "Unit",
    "Total No\nof days", "Bill No", "Advance\nAmt", "Room\nRent", "Licence\nChg",
    "Extra\nBed", "Total Amt",
]
COL_WIDTHS = [6, 14, 9, 20, 16, 11, 14, 8, 9, 10, 10, 10, 8, 11]
BLANK_FILL_COLS = ["B", "C", "G", "I"]  # Army No, Rank, Unit, Bill No — manual fill-in
NUMERIC_TOTAL_COLS = ["J", "K", "L", "M", "N"]  # Advance, Room Rent, Licence, Extra Bed, Total Amt

HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
BLANK_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
TITLE_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True, size=9)
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


async def build_manual_occupancy_workbook(db, month: int, year: int) -> BytesIO:
    import calendar as cal_mod

    days_in_month = cal_mod.monthrange(year, month)[1]
    start_date = f"{year}-{month:02d}-01"
    end_date = f"{year}-{month:02d}-{days_in_month:02d}"
    period_start = date(year, month, 1)
    period_end = date(year, month, days_in_month)

    bookings = await db.bookings.find(
        report_booking_query(start_date, end_date), {"_id": 0}
    ).sort("check_in_date", 1).to_list(5000)

    settings = await db.app_settings.find_one({}, {"_id": 0}) or {}
    all_rooms = await db.rooms.find({}, {"_id": 0}).to_list(100)
    room_maps = build_room_maps(all_rooms)

    rows = []  # one entry per room per booking, in booking order then room order
    for bk in bookings:
        fin = booking_financials(bk, settings, period_start, period_end, room_maps)
        if fin["nights"] == 0 or not fin["rooms"]:
            continue
        is_org = bk.get("is_org", False)
        for room in fin["rooms"]:
            rows.append({
                "name": bk.get("guest_name", ""),
                "mobile_no": bk.get("guest_contact") or "",
                "org_status": "Org" if is_org else "Non-Org",
                "days": room["nights"],
                "advance_amt": bk.get("advance_paid", 0) or 0,
                "room_rent": room["room_rent"],
                "licence_chg": room["license_fee"],
                "extra_bed": (room["revenue"] - room["room_rent"] - room["license_fee"]),
                "total_amt": room["revenue"],
            })

    wb = Workbook()
    ws = wb.active
    ws.title = f"Room Occupancy {cal_mod.month_name[month]} {year}"[:31]

    last_col_letter = get_column_letter(len(HEADERS))

    # Title row
    ws.merge_cells(f"A1:{last_col_letter}1")
    title_cell = ws["A1"]
    title_cell.value = f"SUMMARY OF ROOM OCCUPANCY FOR THE MONTH OF {cal_mod.month_name[month].upper()[:3]} {year}"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20

    # Header row
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.alignment = WRAP_CENTER
        cell.fill = HEADER_FILL
    ws.row_dimensions[2].height = 30

    # Data rows
    r = 3
    for i, row in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=i)  # Ser No
        # B (Army No), C (Rank), G (Unit), I (Bill No) intentionally left blank
        ws.cell(row=r, column=4, value=row["name"])
        ws.cell(row=r, column=5, value=row["mobile_no"])
        ws.cell(row=r, column=6, value=row["org_status"])
        ws.cell(row=r, column=8, value=row["days"])
        ws.cell(row=r, column=10, value=round(row["advance_amt"], 2))
        ws.cell(row=r, column=11, value=round(row["room_rent"], 2))
        ws.cell(row=r, column=12, value=round(row["licence_chg"], 2))
        ws.cell(row=r, column=13, value=round(row["extra_bed"], 2))
        ws.cell(row=r, column=14, value=round(row["total_amt"], 2))
        for col_letter in BLANK_FILL_COLS:
            ws[f"{col_letter}{r}"].fill = BLANK_FILL
        r += 1

    # Total row
    total_row = r
    ws.merge_cells(f"A{total_row}:I{total_row}")
    total_label = ws.cell(row=total_row, column=1, value="Total")
    total_label.font = Font(bold=True)
    total_label.alignment = Alignment(horizontal="center")
    if rows:
        first_data_row, last_data_row = 3, total_row - 1
        for col_letter in NUMERIC_TOTAL_COLS:
            sum_cell = ws[f"{col_letter}{total_row}"]
            sum_cell.value = f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})"
            sum_cell.font = Font(bold=True)
    ws.row_dimensions[total_row].height = 18

    for idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A3"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

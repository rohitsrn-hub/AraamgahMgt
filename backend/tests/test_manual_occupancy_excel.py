"""
Unit test for utils/manual_occupancy_excel.py — verifies the generated
workbook matches the manual ledger's column layout and that its numbers
agree with the shared report_calc financial logic (no independent
recomputation, per the reconciliation fix).
"""
import asyncio
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from utils.manual_occupancy_excel import HEADERS, build_manual_occupancy_workbook

ROOMS = [
    {"id": "r1", "room_number": "C1-01", "category": "Cat I"},
    {"id": "r2", "room_number": "C2-08", "category": "Cat II"},
]


class _FakeCursor:
    def __init__(self, docs):
        self._docs = docs

    def sort(self, *_args, **_kwargs):
        return self

    async def to_list(self, _limit):
        return self._docs


class _FakeCollection:
    def __init__(self, docs):
        self._docs = docs

    def find(self, *_args, **_kwargs):
        return _FakeCursor(self._docs)

    async def find_one(self, *_args, **_kwargs):
        return self._docs[0] if self._docs else None


class _FakeDB:
    def __init__(self, bookings, settings, rooms):
        self.bookings = _FakeCollection(bookings)
        self.app_settings = _FakeCollection([settings] if settings else [])
        self.rooms = _FakeCollection(rooms)


def _booking(**overrides):
    base = {
        "id": "x", "booking_number": "BKTEST", "status": "checked_out",
        "is_org": True, "guest_name": "Test Guest", "guest_contact": "+91 9876543210",
        "check_in_date": "2026-06-10", "check_out_date": "2026-06-11",
        "room_numbers": ["C1-01"], "room_categories": ["Cat I"],
        "extra_beds": 0, "advance_paid": 400,
    }
    base.update(overrides)
    return base


def test_headers_match_manual_layout():
    assert HEADERS[:7] == ["Ser\nNo", "Army No", "Rank", "Name", "Mobile No", "Org /\nNon-Org", "Unit"]
    assert HEADERS[-1] == "Total Amt"


def test_workbook_rows_and_totals():
    bookings = [
        _booking(),
        _booking(booking_number="BKTEST2", is_org=False,
                  room_numbers=["C2-08"], room_categories=["Cat II"],
                  extra_beds=1, guest_name="Second Guest", guest_contact="+91 9123456789"),
    ]
    db = _FakeDB(bookings, {}, ROOMS)

    buf = asyncio.run(build_manual_occupancy_workbook(db, month=6, year=2026))
    wb = load_workbook(buf)
    ws = wb.active

    assert ws["A1"].value == "SUMMARY OF ROOM OCCUPANCY FOR THE MONTH OF JUN 2026"
    assert ws["A2"].value == "Ser\nNo"
    assert ws["D2"].value == "Name"
    assert ws["E2"].value == "Mobile No"
    assert ws["F2"].value == "Org /\nNon-Org"

    # Manual-fill columns (Army No, Rank, Unit, Bill No) stay blank
    assert ws["B3"].value is None
    assert ws["C3"].value is None
    assert ws["G3"].value is None
    assert ws["I3"].value is None

    # Row 1: Cat I org guest, 1 night, no extra bed -> 470 + 30 = 500
    assert ws["D3"].value == "Test Guest"
    assert ws["E3"].value == "+91 9876543210"
    assert ws["F3"].value == "Org"
    assert ws["H3"].value == 1
    assert ws["N3"].value == 500.0

    # Row 2: Cat II non-org guest, 1 night, 1 extra bed -> 570 + 30 + 75 = 675
    assert ws["F4"].value == "Non-Org"
    assert ws["N4"].value == 675.0

    # Total row sums the numeric columns via formula, same convention as the manual sheet
    assert ws["A5"].value == "Total"
    assert ws["N5"].value == "=SUM(N3:N4)"
